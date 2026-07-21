"""
Tests des replis gracieux : chaque intégration externe doit rendre un résultat neutre ("" / False /
None) sans lever d'exception quand sa configuration est absente — le contrat de dégradation
gracieuse commun à tout le projet (conftest.py vide toutes les clés d'API). Teste aussi
l'extraction multi-format des pièces jointes (synthétiques, en mémoire).
"""
import io

import fitz  # PyMuPDF
from docx import Document
from openpyxl import Workbook

import aca.ingestion.attachment_reader as attachment_reader
from aca.agents import enrichment, veille
from aca.integrations import billing, hubspot, notify
from aca.ingestion.attachment_reader import extract_text_from_attachments


# ── notify : aucun canal configuré → False, sans exception ───────────────────────────────────
def test_notify_send_no_channel_returns_false():
    assert notify.send("message de test") is False


# ── hubspot : token absent → inerte ──────────────────────────────────────────────────────────
def test_hubspot_disabled_without_token():
    assert hubspot.is_enabled() is False


def test_hubspot_create_lead_graceful_none():
    result = hubspot.create_lead(
        email_classification="DEVIS",
        extracted_info={"entreprise": "Test SA", "urgence": "haute", "besoin_principal": "10 licences"},
        sender="a@b.fr",
        draft="Bonjour.",
    )
    assert result is None


# ── billing : Stripe non configuré → statistiques seules, jamais d'appel Stripe ──────────────
def test_billing_disabled_without_api_key():
    assert billing.is_enabled() is False


def test_billing_report_usage_returns_stats_without_stripe_call(monkeypatch):
    monkeypatch.setattr(
        billing.stripe, "SubscriptionItem",
        type("Exploding", (), {"create_usage_record": staticmethod(
            lambda *a, **k: (_ for _ in ()).throw(AssertionError("Stripe ne doit pas être appelé"))
        )}),
    )
    stats = billing.report_usage()
    assert set(stats) == {"analyses", "total_entree", "total_sortie", "moyenne_par_analyse"}


# ── enrichment : domaine générique ou Tavily absent → "" ─────────────────────────────────────
def test_enrichment_generic_domain_returns_empty():
    assert enrichment.research_company("quelqu.un@gmail.com") == ""


def test_enrichment_no_tavily_key_returns_empty():
    # Domaine pro mais TAVILY_API_KEY vide (conftest) et cache Sheets indisponible → "".
    assert enrichment.research_company("contact@entreprise-inconnue.fr") == ""


# ── veille : Tavily absent ou requête vide → "" ──────────────────────────────────────────────
def test_veille_no_tavily_key_returns_empty():
    assert veille.search_faq_online("intégrez-vous Salesforce ?") == ""


def test_veille_empty_query_returns_empty():
    assert veille.search_faq_online("   ") == ""


# ── attachment_reader : extraction multi-format + troncature globale ─────────────────────────
def _make_pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    return doc.tobytes()


def _make_docx(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _make_xlsx(cell_value: str) -> bytes:
    workbook = Workbook()
    workbook.active["A1"] = cell_value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_extract_all_three_formats_and_skip_unsupported():
    attachments = [
        ("cahier.pdf", _make_pdf("Contenu PDF cahier des charges")),
        ("annexe.docx", _make_docx("Contenu Word annexe technique")),
        ("budget.xlsx", _make_xlsx("Contenu Excel budget 12000")),
        ("logo.png", b"\x89PNG fake"),          # extension non supportée → ignorée
        ("sans_extension", b"donnees brutes"),  # pas d'extension → ignorée
    ]
    combined = extract_text_from_attachments(attachments)
    assert "Contenu PDF cahier des charges" in combined
    assert "Contenu Word annexe technique" in combined
    assert "Contenu Excel budget 12000" in combined
    # Chaque pièce jointe est préfixée par son nom de fichier pour situer le LLM.
    assert "--- Pièce jointe : cahier.pdf ---" in combined
    assert "logo.png" not in combined


def test_extraction_failure_is_silently_skipped():
    # Un .docx corrompu ne doit ni planter ni polluer la sortie.
    attachments = [
        ("bon.docx", _make_docx("Texte valide")),
        ("corrompu.docx", b"pas un vrai docx"),
    ]
    combined = extract_text_from_attachments(attachments)
    assert "Texte valide" in combined
    assert "corrompu.docx" not in combined


def test_combined_truncation_is_global(monkeypatch):
    # Le budget MAX_CHARS s'applique à l'ENSEMBLE, pas à chaque fichier.
    monkeypatch.setattr(attachment_reader, "MAX_CHARS", 50)
    attachments = [
        ("a.docx", _make_docx("A" * 100)),
        ("b.docx", _make_docx("B" * 100)),
    ]
    combined = extract_text_from_attachments(attachments)
    assert combined.endswith("... [Texte tronqué car trop long]")
    assert len(combined) == 50 + len("... [Texte tronqué car trop long]")


def test_no_attachments_returns_empty():
    assert extract_text_from_attachments([]) == ""
